from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Developer\Documents\Semester_8\Pemograman_API\logistics-api")
TEMPLATE = Path(r"C:\Users\Developer\Documents\Semester_8\Templates_Tugas\Template_Laporan_FIK_UPNVJT_MASTER.docx")
LOGO = Path(r"C:\Users\Developer\Documents\Semester_8\logo_UPN.png")
EVIDENCE = ROOT / "bukti_pengujian"
OUT = ROOT / "output" / "Laporan_UAS_Pemograman_API_Logistics_API_22081010297_FINAL.docx"
FIRST_LINE_CM = 0.27


def clear_body(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_margins(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)


def clear_part(part):
    element = part._element
    for child in list(element):
        element.remove(child)


def clear_headers_footers(doc):
    for section in doc.sections:
        for part in [
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ]:
            clear_part(part)


def tune_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(FIRST_LINE_CM)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size, align in [
        ("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 12, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = align
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(0 if style_name == "Heading 1" else 12)
        style.paragraph_format.space_after = Pt(0 if style_name == "Heading 1" else 6)
        p_pr = style._element.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            p_pr.remove(num_pr)
        if style_name != "Heading 1":
            style.paragraph_format.first_line_indent = None


def p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, bold_label=None):
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.first_line_indent = Cm(FIRST_LINE_CM) if indent else None
    if bold_label and text.startswith(bold_label):
        r = para.add_run(bold_label)
        set_font(r, bold=True)
        r2 = para.add_run(text[len(bold_label):])
        set_font(r2)
    else:
        r = para.add_run(text)
        set_font(r)
    return para


def center(doc, text, size=12, bold=False, before=0, after=0):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    run = para.add_run(text)
    set_font(run, size=size, bold=bold)
    return para


def bab(doc, nomor, judul, first=False):
    if not first:
        doc.add_page_break()
    a = doc.add_paragraph(style="Heading 1")
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.paragraph_format.first_line_indent = None
    a.add_run(f"BAB {nomor}").bold = True
    for run in a.runs:
        set_font(run, size=14, bold=True)
    b = doc.add_paragraph(style="Heading 1")
    b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    b.paragraph_format.first_line_indent = None
    b.paragraph_format.space_after = Pt(18)
    b.add_run(judul.upper()).bold = True
    for run in b.runs:
        set_font(run, size=14, bold=True)


def sub(doc, text):
    para = doc.add_paragraph(style="Heading 2")
    para.paragraph_format.first_line_indent = None
    run = para.add_run(text)
    set_font(run, size=12, bold=True)
    return para


def bullet(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.first_line_indent = Cm(FIRST_LINE_CM)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = para.add_run(text)
    set_font(run)
    return para


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def table_caption(doc, nomor, title):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = None
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(f"Tabel {nomor} {title}")
    set_font(run, bold=True)


def add_table(doc, headers, rows, widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = Cm(widths_cm[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(headers[i])
        set_font(run, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Cm(widths_cm[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = para.add_run(str(value))
            set_font(run)
    p(doc, "", indent=False)
    return table


def image_caption(doc, nomor, title):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = None
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(12)
    run = para.add_run(f"Gambar {nomor} {title}")
    set_font(run, bold=True)


def add_image(doc, filename, caption, nomor, width_in=5.7):
    path = EVIDENCE / filename
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = None
    if path.exists():
        run = para.add_run()
        run.add_picture(str(path), width=Inches(width_in))
    else:
        run = para.add_run(f"[Bukti tidak ditemukan: {filename}]")
        set_font(run, bold=True)
    image_caption(doc, nomor, caption)


def code_table(doc, nomor, title, code):
    table_caption(doc, nomor, title)
    lines = [line.rstrip() for line in code.strip("\n").splitlines()]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.0), Cm(12.8)]
    for i, header in enumerate(["No.", "Potongan Kode"]):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell, top=60, bottom=60, start=80, end=80)
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_font(run, bold=True)
    for idx, line in enumerate(lines, start=1):
        row = table.add_row().cells
        row[0].width = widths[0]
        row[1].width = widths[1]
        for c in row:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c, top=30, bottom=30, start=70, end=70)
        row[0].text = ""
        r0 = row[0].paragraphs[0].add_run(str(idx))
        set_font(r0, name="Courier New", size=9)
        row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row[1].text = ""
        r1 = row[1].paragraphs[0].add_run(line[:120])
        set_font(r1, name="Courier New", size=9)
    p(doc, "", indent=False)


def cover(doc):
    center(doc, "LAPORAN UJIAN AKHIR SEMESTER", size=14, bold=True, before=36)
    center(doc, "PENGEMBANGAN WEBSITE PENYEDIA API LOGISTIK", size=14, bold=True)
    center(doc, "MENGGUNAKAN FASTAPI DAN API KEY", size=14, bold=True, after=24)
    if LOGO.exists():
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = None
        para.add_run().add_picture(str(LOGO), width=Inches(1.85))
    else:
        center(doc, "[LOGO UPN]", bold=True)
    center(doc, "Mata Kuliah :", after=0)
    center(doc, "Pemograman API", after=12)
    center(doc, "Dosen Pengampu:", after=0)
    center(doc, "Muhammad Muharrom Al Haromainy, S.Kom., M.Kom.", after=12)
    center(doc, "Disusun Oleh:", after=0)
    center(doc, "Yogi Prasetyo", after=0)
    center(doc, "22081010297", after=24)
    center(doc, "PROGRAM STUDI INFORMATIKA", bold=True, before=60)
    center(doc, "FAKULTAS ILMU KOMPUTER", bold=True)
    center(doc, 'UNIVERSITAS PEMBANGUNAN NASIONAL "VETERAN" JAWA TIMUR', bold=True)
    center(doc, "2025/2026", bold=True)
    doc.add_page_break()


def build():
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    for section in doc.sections:
        set_margins(section)
    clear_headers_footers(doc)
    tune_styles(doc)

    cover(doc)

    bab(doc, "I", "Pendahuluan", first=True)
    sub(doc, "1.1 Latar Belakang")
    for text in [
        "Perkembangan aplikasi modern membuat API menjadi komponen penting dalam pertukaran data antar sistem. API memungkinkan sebuah layanan dipakai oleh aplikasi lain tanpa membuka akses langsung ke database. Pada MK Pemograman API, pemahaman tersebut perlu dibuktikan melalui proyek yang berjalan dan dapat diuji.",
        "Soal UAS meminta mahasiswa membangun website penyedia API berbasis database lokal, menyediakan login dan daftar akun, membuat API Key, menerapkan autentikasi, menyediakan proses View Data, Create, Update, dan Delete Data, serta menguji API menggunakan Postman.",
        "Penulis mengembangkan ShipStream - Logistics Rate & Tracking API. Sistem ini mengangkat domain logistik karena memiliki kebutuhan API yang jelas, seperti data kota, data kurir, perhitungan ongkos kirim, pembuatan resi, tracking paket, dan dashboard client."
    ]:
        p(doc, text)
    sub(doc, "1.2 Rumusan Masalah")
    p(doc, "Bagaimana membangun website penyedia API logistik berbasis database lokal yang memiliki register, login, API Key, autentikasi JWT/API Key, operasi data, dokumentasi, web client, deployment, dan bukti pengujian Postman sesuai soal UAS Pemograman API?")
    sub(doc, "1.3 Tujuan")
    for item in [
        "Membangun backend API menggunakan FastAPI untuk kebutuhan layanan logistik.",
        "Menerapkan register, OTP, login JWT, dan manajemen API Key.",
        "Menyediakan endpoint data logistik yang dapat digunakan oleh web client.",
        "Menyediakan dokumentasi Swagger/OpenAPI, API guide, dan Postman collection.",
        "Menguji endpoint menggunakan pytest dan Postman."
    ]:
        bullet(doc, item)
    sub(doc, "1.4 Manfaat")
    p(doc, "Laporan ini bermanfaat sebagai dokumentasi proses pengembangan API dari analisis requirement, implementasi endpoint, pembuatan dokumentasi, pengujian, sampai persiapan deployment. Bagi penulis, proyek ini memperkuat pemahaman tentang backend API yang aman dan dapat digunakan oleh client.")

    bab(doc, "II", "Dasar Teori")
    sub(doc, "2.1 REST API")
    p(doc, "REST API adalah gaya arsitektur web service yang memanfaatkan HTTP method seperti GET, POST, PUT, PATCH, dan DELETE untuk mengelola resource. Response API biasanya dikirim dalam format JSON sehingga mudah digunakan oleh aplikasi web, mobile, atau service lain.")
    sub(doc, "2.2 FastAPI")
    p(doc, "FastAPI adalah framework Python modern untuk membangun API. Framework ini mendukung validasi schema melalui Pydantic, dependency injection, middleware, dan dokumentasi otomatis berbasis OpenAPI. Pada proyek ini FastAPI dipakai karena cepat, ringkas, dan cocok untuk API yang perlu banyak validasi request.")
    sub(doc, "2.3 JWT dan API Key")
    p(doc, "JWT digunakan untuk membuktikan identitas user setelah login. Token dikirim melalui header Authorization. API Key digunakan untuk mengakses endpoint logistik melalui header X-API-Key. Pemisahan ini membuat alur lebih jelas: JWT untuk manajemen akun, API Key untuk konsumsi layanan API.")
    sub(doc, "2.4 SQLAlchemy dan Database Lokal")
    p(doc, "SQLAlchemy digunakan sebagai ORM untuk memetakan tabel database ke objek Python. Database lokal dipakai untuk menyimpan data user, API key, kota, kurir, tracking, transaksi, dan data pendukung lain.")
    sub(doc, "2.5 Pengujian API")
    p(doc, "Pengujian API dilakukan dengan dua cara. Pytest digunakan untuk menguji fungsi dan route secara otomatis, sedangkan Postman digunakan untuk membuktikan alur pemakaian API dari sudut pandang pengguna.")

    bab(doc, "III", "Metodologi")
    sub(doc, "3.1 Spesifikasi Alat dan Bahan")
    table_caption(doc, "3.1", "Spesifikasi Alat dan Bahan")
    add_table(doc, ["Komponen", "Keterangan"], [
        ["Bahasa pemrograman", "Python 3.13"],
        ["Framework backend", "FastAPI"],
        ["Database lokal", "SQLite melalui SQLAlchemy"],
        ["Autentikasi", "JWT Bearer Token dan X-API-Key"],
        ["Pengujian", "Pytest dan Postman"],
        ["Dokumentasi", "Swagger UI, ReDoc, docs/API_GUIDE.md, docs/ENDPOINTS.md"],
        ["Web client", "frontend/index.html dan frontend/dashboard.html"],
        ["Deployment", "Railway untuk backend dan web client yang disajikan dari aplikasi FastAPI"]
    ], [4.0, 9.8])
    doc.add_page_break()
    sub(doc, "3.2 Analisis Requirement Soal")
    table_caption(doc, "3.2", "Kesesuaian Requirement UAS dengan Implementasi")
    add_table(doc, ["No.", "Requirement Soal", "Implementasi pada Proyek", "Status"], [
        ["1", "Website penyedia API dengan basis data lokal", "FastAPI dengan database SQLite/SQLAlchemy dan endpoint dokumentasi", "Terpenuhi"],
        ["2", "API Key untuk menggunakan API", "Endpoint /apikey/generate, /apikey/list, /apikey/revoke/{key_id}", "Terpenuhi"],
        ["3", "Login dan daftar akun sederhana", "Endpoint /auth/register, /auth/verify-code, /auth/login", "Terpenuhi"],
        ["4", "Autentikasi JWT/API Key", "JWT untuk akun dan X-API-Key untuk endpoint /v1", "Terpenuhi"],
        ["5", "View, Create, Update, Delete data", "Daftar kota/kurir, create shipment, update tracking/API key, revoke API key", "Terpenuhi"],
        ["6", "Implementasi dengan Postman", "Collection postman/collection.json dan screenshot bukti_pengujian", "Terpenuhi"],
        ["7", "Landing page dan dokumentasi API", "Root URL https://web-production-24009.up.railway.app/, /docs, /redoc, docs/API_GUIDE.md", "Terpenuhi"],
        ["8", "Website client pengguna API", "frontend/dashboard.html sebagai dashboard client", "Terpenuhi"],
        ["9", "Deploy API dan client", "Aplikasi berjalan pada Railway melalui https://web-production-24009.up.railway.app/", "Terpenuhi"]
    ], [1.0, 4.0, 6.6, 2.2])
    sub(doc, "3.3 Alur Kerja")
    for item in [
        "Membaca dan memetakan requirement dari file soal EAS API 2526.pdf.",
        "Mendesain struktur aplikasi FastAPI dalam folder app, routers, schemas, models, dan utils.",
        "Mengimplementasikan autentikasi akun, OTP, JWT, API Key, dan validasi request.",
        "Mengimplementasikan endpoint logistik untuk kota, kurir, ongkos kirim, shipment, tracking, label PDF, dan quota.",
        "Membuat dokumentasi API serta web client agar API dapat digunakan dari browser.",
        "Menjalankan pengujian otomatis dengan pytest dan pengujian manual dengan Postman."
    ]:
        bullet(doc, item)

    bab(doc, "IV", "Hasil dan Pembahasan")
    sub(doc, "4.1 Ringkasan Implementasi Sistem")
    p(doc, "Aplikasi yang dihasilkan bernama ShipStream - Logistics Rate & Tracking API. Backend menyediakan endpoint autentikasi, API Key, billing, profil, telegram webhook, dan data logistik. Web client disajikan melalui route root dan dashboard sehingga pengguna dapat mengakses layanan API dari browser.")
    table_caption(doc, "4.1", "Kelompok Endpoint Utama")
    add_table(doc, ["Kelompok", "Endpoint Contoh", "Fungsi"], [
        ["Auth", "/auth/register, /auth/login, /auth/verify-code", "Registrasi, verifikasi OTP, dan login JWT"],
        ["API Key", "/apikey/generate, /apikey/list", "Membuat dan mengelola API Key"],
        ["Logistik", "/v1/couriers, /v1/cities, /v1/cost", "Mengambil data logistik dan menghitung ongkos kirim"],
        ["Shipment", "/v1/shipments, /v1/tracking/{awb}", "Membuat resi dan melacak status pengiriman"],
        ["Dokumentasi", "/docs, /redoc", "Dokumentasi API interaktif"]
    ], [3.0, 5.0, 5.8])
    doc.add_page_break()
    table_caption(doc, "4.2", "Informasi Akses Deployment dan Akun Tester")
    add_table(doc, ["Item", "Keterangan"], [
        ["Root URL / Landing Page", "https://web-production-24009.up.railway.app/"],
        ["Dokumentasi Swagger", "https://web-production-24009.up.railway.app/docs"],
        ["Akun tester", "Username: testuser04"],
        ["Password tester", "TestPassword123!"]
    ], [4.0, 9.8])
    sub(doc, "4.2 Dokumentasi Kode Program")
    p(doc, "Potongan kode berikut menunjukkan bagian inti yang membuktikan website penyedia API, routing, autentikasi API Key, dan logika ongkos kirim.")
    doc.add_page_break()
    code_table(doc, "4.3", "Inisialisasi FastAPI dan Router Utama", """
app = FastAPI(
    title="Logistics Rate & Tracking API",
    description="API for checking shipping rates and tracking packages (EAS Semester 8)",
    version="1.0.0"
)
app.include_router(auth.router)
app.include_router(apikey.router)
app.include_router(logistics.router)
app.include_router(billing.router)
app.include_router(profile.router)
""")
    code_table(doc, "4.4", "Endpoint Generate API Key", """
@router.post("/generate", response_model=schemas.APIKey)
def generate_key(key_data: schemas.APIKeyCreate,
                 current_user: models.User = Depends(get_current_user),
                 db: Session = Depends(get_db)):
    key = generate_api_key()
    db_key = models.APIKey(
        key=key,
        label=key_data.label,
        owner_id=current_user.id,
        request_limit=key_data.request_limit
    )
""")
    doc.add_page_break()
    code_table(doc, "4.5", "Endpoint Perhitungan Ongkos Kirim", """
@router.post("/cost", summary="Kalkulasi Ongkos Kirim (Scale-up)")
@limiter.limit("30/minute")
async def calculate_cost(request: Request,
                         payload: schemas.CostRequest,
                         response: Response,
                         db: Session = Depends(get_db),
                         api_key=Depends(validate_api_key)):
    origin = db.query(models.City).filter(models.City.id == payload.origin).first()
    destination = db.query(models.City).filter(models.City.id == payload.destination).first()
""")
    sub(doc, "4.3 Hasil Pengujian Otomatis")
    p(doc, "Pengujian otomatis dijalankan ulang menggunakan perintah python -m pytest -q. Hasil terbaru menunjukkan seluruh test lulus, yaitu 67 passed. Warning yang muncul berupa deprecation warning dari dependensi dan tidak menggagalkan test.")
    table_caption(doc, "4.6", "Ringkasan Hasil Pytest")
    add_table(doc, ["Item", "Hasil"], [
        ["Total test", "67"],
        ["Passed", "67"],
        ["Failed", "0"],
        ["Status", "Layak digunakan untuk demonstrasi UAS"],
        ["Catatan", "Terdapat 144 warning deprecation, namun tidak memengaruhi kelulusan test"]
    ], [4.0, 9.8])
    sub(doc, "4.4 Hasil Pengujian Manual Postman")
    p(doc, "Pengujian manual dilakukan secara berurutan. JWT yang diperoleh dari login digunakan untuk membuat API Key, kemudian API Key dipakai pada header X-API-Key untuk endpoint logistik.")
    images = [
        ("bukti register postman.png", "Register akun melalui Postman"),
        ("bukti ambil otp postman.png", "Pengambilan OTP debug"),
        ("bukti verifikasi otp postman.png", "Verifikasi OTP berhasil"),
        ("bukti login postman.png", "Login berhasil dan JWT diterbitkan"),
        ("bukti create api key postman.png", "Pembuatan API Key berhasil"),
        ("bukti ambil data courier postman.png", "Pengambilan data kurir"),
        ("bukti ambil data kota postman.png", "Pengambilan data kota"),
        ("bukti ambil data cost postman.png", "Pengujian perhitungan ongkos kirim")
    ]
    for idx, (filename, caption) in enumerate(images, start=1):
        add_image(doc, filename, caption, f"4.{idx}", width_in=5.65)
    sub(doc, "4.5 Analisis Hasil")
    for item in [
        "Alur register, OTP, verifikasi, login, dan penerbitan JWT berjalan sesuai skenario.",
        "Endpoint generate API Key berhasil dipakai setelah pengguna login.",
        "Endpoint logistik dapat diakses menggunakan header X-API-Key, sehingga requirement API Key terpenuhi.",
        "Data kurir, data kota, dan perhitungan ongkos kirim berhasil diuji melalui Postman.",
        "Hasil pytest 67 passed memperkuat bahwa fungsi utama sistem stabil untuk kebutuhan pengumpulan UAS."
    ]:
        bullet(doc, item)

    bab(doc, "V", "Penutup")
    sub(doc, "5.1 Kesimpulan")
    p(doc, "ShipStream - Logistics Rate & Tracking API berhasil dikembangkan sesuai requirement UAS Pemograman API. Sistem menyediakan website penyedia API, database lokal, register dan login, OTP, JWT, API Key, endpoint logistik, dokumentasi API, web client, serta bukti pengujian.")
    p(doc, "Pengujian otomatis menghasilkan 67 passed. Pengujian manual melalui Postman juga menunjukkan alur penggunaan API berjalan mulai dari register sampai perhitungan ongkos kirim.")
    sub(doc, "5.2 Saran")
    for item in [
        "Gunakan PostgreSQL permanen pada deployment agar data tidak hilang setelah redeploy.",
        "Matikan endpoint debug OTP saat aplikasi digunakan di lingkungan production.",
        "Gunakan Redis permanen untuk rate limiting bila trafik meningkat.",
        "Tambahkan monitoring log agar error deployment lebih mudah dilacak.",
        "Perluas pengujian integrasi ke RajaOngkir ketika API key eksternal tersedia."
    ]:
        bullet(doc, item)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
