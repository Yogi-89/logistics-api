from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output" / "Laporan_Pemograman_API_Pengembangan_REST_API_Logistik_22081010297.docx"
OUT = ROOT / "output" / "Laporan_Pemograman_API_REST_API_Logistik_22081010297_UPDATED.docx"
EVIDENCE = ROOT / "bukti_pengujian"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if "Table Grid" in [s.name for s in doc.styles]:
        table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True)
        set_cell_shading(hdr[i], "E8EEF5")
        hdr[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


def add_paragraph(doc, text="", style=None, bold_label=None):
    p = doc.add_paragraph(style=style)
    if bold_label and text.startswith(bold_label):
        r = p.add_run(bold_label)
        r.bold = True
        p.add_run(text[len(bold_label):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text):
    style_names = [s.name for s in doc.styles]
    p = doc.add_paragraph(style="List Paragraph" if "List Paragraph" in style_names else None)
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    return p


def add_picture(doc, filename, caption):
    image_path = EVIDENCE / filename
    if not image_path.exists():
        add_paragraph(doc, f"[Bukti tidak ditemukan: {filename}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(5.8))
    add_caption(doc, caption)


def delete_from_heading(doc, heading_text):
    start = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == heading_text:
            start = idx
            break
    if start is None:
        raise RuntimeError(f"Heading not found: {heading_text}")
    for p in list(doc.paragraphs[start:]):
        p._element.getparent().remove(p._element)


def tune_styles(doc):
    if "Normal" in [s.name for s in doc.styles]:
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.1

    for name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        if name not in [s.name for s in doc.styles]:
            continue
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)


def main():
    doc = Document(str(SOURCE))
    tune_styles(doc)
    delete_from_heading(doc, "BAB IV")

    doc.add_page_break()
    h = doc.add_paragraph("BAB IV", style="Heading 1")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h2 = doc.add_paragraph("HASIL PENGUJIAN DAN PEMBAHASAN", style="Heading 1")
    h2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_paragraph(doc, "4.1 Lingkungan Pengujian", style="Heading 2")
    add_paragraph(
        doc,
        "Pengujian akhir dilakukan pada deployment Railway dengan base URL "
        "https://web-production-24009.up.railway.app. Pengujian menggunakan Postman "
        "untuk memvalidasi alur autentikasi, pembuatan API key, serta endpoint data logistik. "
        "Selain pengujian manual, pengujian otomatis juga dijalankan menggunakan pytest dan seluruh "
        "test case berhasil lulus."
    )
    add_table(
        doc,
        ["Komponen", "Keterangan"],
        [
            ["Framework", "FastAPI"],
            ["Database deployment", "SQLite fallback pada Railway untuk kebutuhan demo"],
            ["Autentikasi", "JWT untuk login dan API key untuk endpoint logistik"],
            ["Rate limiter", "REDIS_URL=memory:// untuk mode demo"],
            ["Integrasi ongkir", "RajaOngkir/Komerce melalui RAJA_ONGKIR_API_KEY"],
            ["Deployment", "Railway, start command python railway_start.py"],
            ["Pengujian otomatis", "67 passed"],
        ],
        [2.0, 4.3],
    )

    add_paragraph(doc, "4.2 Ringkasan Alur Pengujian", style="Heading 2")
    add_table(
        doc,
        ["No.", "Endpoint", "Metode", "Tujuan", "Hasil"],
        [
            ["1", "/auth/register", "POST", "Registrasi akun pengguna", "Berhasil, user dibuat"],
            ["2", "/auth/debug/otp/{username}", "GET", "Mengambil OTP mode debug", "Berhasil, OTP tersedia"],
            ["3", "/auth/verify-code", "POST", "Verifikasi kode OTP", "Berhasil, akun verified"],
            ["4", "/auth/login", "POST", "Login dengan header X-Captcha-Token", "Berhasil, JWT diterbitkan"],
            ["5", "/apikey/generate", "POST", "Membuat API key pengguna", "Berhasil, API key aktif"],
            ["6", "/v1/couriers", "GET", "Melihat daftar kurir", "Berhasil, data kurir tampil"],
            ["7", "/v1/cities", "GET", "Melihat daftar kota", "Berhasil, data kota tampil"],
            ["8", "/v1/cost", "POST", "Menghitung ongkos kirim", "Berhasil setelah rajaongkir_id di-seed"],
        ],
        [0.45, 1.55, 0.8, 2.2, 1.4],
    )

    add_paragraph(doc, "4.3 Hasil Pengujian Manual Postman", style="Heading 2")
    add_paragraph(
        doc,
        "Pengujian manual dilakukan secara berurutan agar alur dependensi antar-endpoint terlihat jelas. "
        "Endpoint login menghasilkan token JWT, kemudian token tersebut dipakai untuk membuat API key. "
        "API key selanjutnya digunakan sebagai header X-API-Key pada endpoint logistik."
    )

    pictures = [
        ("bukti register postman.png", "Gambar 4.1 Pengujian register akun melalui Postman"),
        ("bukti ambil otp postman.png", "Gambar 4.2 Pengambilan OTP debug untuk akun uji"),
        ("bukti verifikasi otp postman.png", "Gambar 4.3 Verifikasi OTP berhasil"),
        ("bukti login postman.png", "Gambar 4.4 Login berhasil dan access token diterbitkan"),
        ("bukti create api key postman.png", "Gambar 4.5 Pembuatan API key berhasil"),
        ("bukti ambil data courier postman.png", "Gambar 4.6 Pengambilan data kurir"),
        ("bukti ambil data kota postman.png", "Gambar 4.7 Pengambilan data kota"),
        ("bukti ambil data cost postman.png", "Gambar 4.8 Pengujian perhitungan ongkos kirim"),
    ]
    for filename, caption in pictures:
        add_picture(doc, filename, caption)

    add_paragraph(doc, "4.4 Temuan dan Perbaikan Saat Deployment", style="Heading 2")
    add_paragraph(
        doc,
        "Pada proses deployment ditemukan beberapa masalah konfigurasi yang memengaruhi runtime. "
        "Masalah tersebut telah diperbaiki agar aplikasi dapat berjalan stabil pada Railway."
    )
    add_table(
        doc,
        ["Masalah", "Dampak", "Perbaikan"],
        [
            [
                "DATABASE_URL berisi placeholder",
                "Aplikasi gagal start karena SQLAlchemy tidak dapat membaca URL database",
                "Menambahkan sanitasi environment dan fallback ke SQLite",
            ],
            [
                "REDIS_URL berisi placeholder",
                "Rate limiter gagal import sehingga healthcheck Railway gagal",
                "Menggunakan REDIS_URL=memory:// untuk demo",
            ],
            [
                "Tabel cities dan couriers kosong setelah redeploy",
                "Endpoint /v1/cities dan /v1/couriers mengembalikan array kosong",
                "Menambahkan seed data referensi otomatis pada railway_start.py",
            ],
            [
                "rajaongkir_id kota belum terisi",
                "Endpoint /v1/cost menolak request karena ID tujuan belum ada",
                "Menambahkan seed RajaOngkir ID otomatis berbasis SQLAlchemy",
            ],
        ],
        [1.8, 2.25, 2.25],
    )

    add_paragraph(doc, "4.5 Analisis Hasil", style="Heading 2")
    add_bullet(doc, "Alur autentikasi berjalan lengkap: register, OTP, verifikasi, login, dan penerbitan JWT.")
    add_bullet(doc, "API key berhasil dibuat dan digunakan untuk mengakses endpoint logistik.")
    add_bullet(doc, "Data referensi kota dan kurir tersedia setelah mekanisme seed otomatis ditambahkan.")
    add_bullet(doc, "Endpoint ongkos kirim dapat memvalidasi kota, kurir, berat, dan menggunakan fallback bila integrasi eksternal tidak tersedia.")
    add_bullet(doc, "Pengujian otomatis menghasilkan 67 passed sehingga fungsi utama aplikasi dinilai layak untuk demonstrasi dan pengumpulan tugas.")

    doc.add_page_break()
    h = doc.add_paragraph("BAB V", style="Heading 1")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h2 = doc.add_paragraph("KESIMPULAN DAN SARAN", style="Heading 1")
    h2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_paragraph(doc, "5.1 Kesimpulan", style="Heading 2")
    add_paragraph(
        doc,
        "Berdasarkan implementasi dan pengujian, aplikasi Logistics Rate & Tracking API berhasil memenuhi "
        "kebutuhan utama tugas Pemrograman API. Sistem telah menyediakan registrasi akun, login, OTP, "
        "JWT, API key, dokumentasi Swagger, serta endpoint logistik untuk data kota, data kurir, tracking, "
        "dan perhitungan ongkos kirim. Deployment pada Railway berhasil dilakukan dan bukti pengujian "
        "manual melalui Postman telah terdokumentasi."
    )
    add_paragraph(
        doc,
        "Pengujian otomatis menggunakan pytest menunjukkan hasil 67 passed. Hasil tersebut memperkuat "
        "bahwa fungsi utama aplikasi berjalan sesuai rancangan. Beberapa kendala deployment terkait "
        "environment variable dan data seed berhasil diperbaiki sehingga aplikasi lebih siap digunakan "
        "untuk demonstrasi."
    )

    add_paragraph(doc, "5.2 Saran", style="Heading 2")
    add_bullet(doc, "Gunakan PostgreSQL Railway agar data akun dan API key tidak hilang setelah redeploy.")
    add_bullet(doc, "Nonaktifkan DEBUG_MODE setelah pengujian OTP selesai agar endpoint debug tidak terbuka.")
    add_bullet(doc, "Gunakan Redis permanen bila aplikasi dipakai untuk trafik nyata, bukan hanya mode demo.")
    add_bullet(doc, "Tambahkan pengujian integrasi khusus untuk RajaOngkir bila kuota API eksternal tersedia.")
    add_bullet(doc, "Tambahkan role admin jika sistem akan dipakai lebih lanjut untuk manajemen pengguna.")

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
